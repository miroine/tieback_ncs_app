import sys, io, zipfile, copy, yaml
import tb_catalog as c, tb_network as n, tb_cost, tb_schedule, tb_flowassurance as fa, tb_report as rep
from _harness import Suite
S = Suite("test_report")
FIX = yaml.safe_load(open("test_fixtures/demo_field_a_tieback.yaml"))
def demo(): return n.Layout.from_dict(copy.deepcopy(FIX), c.Catalog())
DATA = rep.build_report("Field A tie-back", demo(), tb_cost.CostSettings(), tb_schedule.ScheduleSettings(),
                        fa.FASettings(), author="Concept Development")

def is_docx():
    z = zipfile.ZipFile(io.BytesIO(DATA))
    names = z.namelist()
    assert "word/document.xml" in names and "[Content_Types].xml" in names
    assert len(DATA) > 20000
S.check("produces a valid .docx package", is_docx)
def text():
    from docx import Document
    doc = Document(io.BytesIO(DATA))
    body = "\n".join(p.text for p in doc.paragraphs)
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    for section in ("Concept summary", "Design checks", "Equipment and quantities", "Cost estimate",
                    "Schedule", "Flow assurance", "Basis and limitations"):
        assert section in heads, section
    assert "Field A tie-back" in body and "Concept Development" in body
    assert "not a substitute for OLGA" in body and "Drilling and completion cost is excluded" in body
S.check("all sections, the title, author and the limitations text are present", text)
def tables_and_figures():
    from docx import Document
    doc = Document(io.BytesIO(DATA))
    assert len(doc.tables) >= 6
    summary = {r.cells[0].text: r.cells[1].text for r in doc.tables[0].rows}
    assert summary["Wells"] == "4" and "MUSD" in summary["CAPEX incl. contingency"]
    z = zipfile.ZipFile(io.BytesIO(DATA))
    images = [x for x in z.namelist() if x.startswith("word/media/")]
    assert len(images) >= 3          # CAPEX, annual profile, route section
S.check("summary table is populated and charts are embedded", tables_and_figures)
def without_fa():
    data = rep.build_report("No FA", demo(), tb_cost.CostSettings(), tb_schedule.ScheduleSettings(),
                            include_flow_assurance=False)
    from docx import Document
    doc = Document(io.BytesIO(data))
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "Flow assurance" not in heads and "Cost estimate" in heads
S.check("flow assurance can be left out", without_fa)
def with_findings():
    lay = demo(); lay.edges["FL1"].diameter_in = 0        # provokes a design error
    data = rep.build_report("Broken", lay, tb_cost.CostSettings(), tb_schedule.ScheduleSettings(), fa.FASettings())
    from docx import Document
    doc = Document(io.BytesIO(data))
    txt = "\n".join(cell.text for t in doc.tables for r in t.rows for cell in r.cells)
    assert "Error" in txt and "diameter" in txt.lower()
S.check("design findings are reported in the document", with_findings)
def empty_layout():
    data = rep.build_report("Empty", n.Layout(c.Catalog()), tb_cost.CostSettings(), tb_schedule.ScheduleSettings())
    assert len(data) > 10000
S.check("an empty layout still produces a document", empty_layout)
sys.exit(0 if S.report() else 1)
