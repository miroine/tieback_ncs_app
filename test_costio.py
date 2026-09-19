import sys, io as _io
import tb_catalog as c, tb_costio as cio
from _harness import Suite
S = Suite("test_costio")
CAT = c.Catalog()
WB = cio.catalog_to_workbook(CAT)

S.check("workbook round-trip reproduces the catalog exactly",
        lambda: cio.catalog_from_workbook(WB).to_dict() == CAT.to_dict())
S.check("CSV round-trip reproduces the catalog exactly",
        lambda: (lambda t: cio.catalog_from_csv(t[0].encode(), t[1].encode()).to_dict() == CAT.to_dict())(cio.catalog_to_csv(CAT)))
def sheets():
    from openpyxl import load_workbook
    wb = load_workbook(_io.BytesIO(WB))
    assert wb.sheetnames == ["items", "spreads", "notes"]
    assert wb["items"].max_row == len(CAT.items) + 1 and wb["spreads"].max_row == len(CAT.spreads) + 1
    assert wb["items"].cell(row=1, column=1).value == "item_id"
S.check("workbook has items, spreads and notes sheets with headers", sheets)
def edited():
    items, spreads = cio.catalog_to_csv(CAT)
    edited_csv = items.replace("5500000.0", "6100000.0")
    cat2 = cio.catalog_from_csv(edited_csv.encode(), spreads.encode())
    assert cat2.get("xt_vxt_10k").procurement_usd == 6100000.0
S.check("edited rate is picked up", edited)
def csv_only_items():
    items, _ = cio.catalog_to_csv(CAT)
    cat2 = cio.catalog_from_csv(items.encode())
    assert len(cat2.spreads) == len(c.Catalog().spreads)      # defaults kept when no spreads file
S.check("items-only CSV keeps the default vessel spreads", csv_only_items)
def missing_cols():
    csv_txt = "item_id,name,category,procurement_usd\nplet_x,Custom PLET,plet,2500000\n"
    cat2 = cio.catalog_from_csv(csv_txt.encode())
    it = cat2.get("plet_x")
    assert it.procurement_usd == 2500000 and it.cost_basis == "unit" and it.uncertainty == (0.9, 1.0, 1.3)
S.check("missing columns fall back to defaults", missing_cols)
def blank_and_extra():
    csv_txt = ("item_id,name,category,procurement_usd,my_notes\n"
               "plet_x,Custom PLET,plet,2500000,ignore me\n,,,,\n")
    assert len(cio.catalog_from_csv(csv_txt.encode()).items) == 1
S.check("blank rows skipped and unknown columns ignored", blank_and_extra)
S.raises("bad number names the row", ValueError,
         lambda: cio.catalog_from_csv(b"item_id,name,category,procurement_usd\nx,X,plet,not-a-number\n"))
S.raises("invalid category names the row", ValueError,
         lambda: cio.catalog_from_csv(b"item_id,name,category\nx,X,spaceship\n"))
S.raises("empty sheet raises", ValueError, lambda: cio.catalog_from_csv(b"item_id,name,category\n"))
def no_items_sheet():
    from openpyxl import Workbook
    wb = Workbook(); wb.active.title = "rates"; buf = _io.BytesIO(); wb.save(buf)
    cio.catalog_from_workbook(buf.getvalue())
S.raises("workbook without an items sheet raises", ValueError, no_items_sheet)
def library_file():
    data = open("library/cost_library_template.xlsx", "rb").read()
    assert cio.catalog_from_workbook(data).to_dict() == CAT.to_dict()
S.check("shipped Excel template reloads into an identical catalog", library_file)
sys.exit(0 if S.report() else 1)
